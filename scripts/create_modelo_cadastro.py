from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from pathlib import Path

headers = [
    'unidade',
    'nome',
    'nome_social',
    'data_nascimento',
    'cpf',
    'rg',
    'orgao_rg',
    'nacionalidade',
    'natural_uf',
    'natural_cidade',
    'nome_mae',
    'cpf_mae',
    'nome_pai',
    'cpf_pai',
    'responsavel_tipo',
    'responsavel_nome',
    'responsavel_cpf',
    'telefone_resp',
    'cep',
    'rua',
    'numero',
    'bairro',
    'cidade',
    'uf',
    'whatsapp',
    'email',
    'nivel',
    'genero',
    'raca_cor',
    'renda_familiar',
    'residente_maior_renda',
    'pessoas_residencia',
    'ocupacao',
    'beneficio_social_status',
    'beneficio_social_nome',
    'meio_transporte',
    'saude_laudo',
    'saude_medicacao',
    'saude_medicamento_nome',
    'saude_observacoes',
    'autorizacao_imagem',
    'turmas_selecionadas',
]

wb = Workbook()
ws = wb.active
ws.title = 'Cadastro'
ws.append(headers)

header_fill = PatternFill(fill_type='solid', fgColor='1F4E78')
header_font = Font(bold=True, color='FFFFFF')
for cell in ws[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal='center', vertical='center')

for col_idx in range(1, len(headers) + 1):
    ws.column_dimensions[get_column_letter(col_idx)].width = 22

ws.freeze_panes = 'A2'

ws2 = wb.create_sheet('Instrucoes')
ws2.append(['Campo', 'Orientação de preenchimento'])
for row in [
    ['unidade', 'Nome exato da unidade cadastrada no sistema (ex.: MGB)'],
    ['nome', 'Nome completo do aluno'],
    ['nome_social', 'Nome social, se houver'],
    ['data_nascimento', 'Formato YYYY-MM-DD ou DD/MM/YYYY'],
    ['cpf', 'Somente números ou com máscara'],
    ['genero', 'Ex.: Masculino, Feminino, Não binário'],
    ['raca_cor', 'Ex.: Branco, Pardo, Preto, Indígena, Amarelo'],
    ['saude_laudo', 'Sim ou Não'],
    ['autorizacao_imagem', 'Sim ou Não'],
    ['turmas_selecionadas', 'IDs das turmas separadas por vírgula'],
]:
    ws2.append(row)

for cell in ws2[1]:
    cell.fill = header_fill
    cell.font = header_font

for col_idx in range(1, 3):
    ws2.column_dimensions[get_column_letter(col_idx)].width = 30

path = Path('scripts/modelo_cadastro_alunos.xlsx')
wb.save(path)
print(path)
print('headers', len(headers))
